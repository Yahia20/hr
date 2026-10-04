from io import BytesIO
from datetime import date

from docx import Document
from pypdf import PdfWriter

from app.ats.extract import extract_text
from app.ats.parse import parse_cv


EN_CV = """Nora Example
nora@example.test
Mobile: +966 50 000 0001
Professional Experience
Reservation Agent | January 2019 - Present
Used Amadeus, PNR, fare rules and Excel for customer service.
Education
Bachelor of Tourism
Languages
Arabic, English
"""

AR_CV = """ليلى مثال
layla@example.test
هاتف: +20 10 0000 0001
الخبرة المهنية
موظفة حجوزات | يناير 2020 - حتى الآن
استخدام أماديوس وحجز الفنادق وخدمة العملاء.
المؤهلات العلمية
بكالوريوس السياحة
اللغات
العربية والإنجليزية
"""


def test_parse_english_cv_fields_and_skill_evidence():
    parsed = parse_cv(EN_CV)
    assert parsed.name == "Nora Example"
    assert parsed.email == "nora@example.test"
    assert parsed.phone == "+966500000001"
    assert parsed.education == "bachelor"
    assert parsed.years_experience is not None and parsed.years_experience >= 6
    assert {"Arabic", "English"} <= set(parsed.languages)
    assert {"amadeus", "pnr", "fare_rules", "excel", "customer_service"} <= {skill.id for skill in parsed.skills}
    assert all(skill.evidence and all(item.line for item in skill.evidence) for skill in parsed.skills)


def test_parse_arabic_cv_fields_dates_and_sections():
    parsed = parse_cv(AR_CV)
    assert parsed.name == "ليلى مثال"
    assert parsed.email == "layla@example.test"
    assert parsed.phone == "+201000000001"
    assert parsed.education == "bachelor"
    assert parsed.years_experience is not None and parsed.years_experience >= 5
    assert {"Arabic", "English"} <= set(parsed.languages)
    assert {"amadeus", "hotel_booking", "customer_service"} <= {skill.id for skill in parsed.skills}
    assert "experience" in parsed.sections


def test_overlapping_date_ranges_are_merged_not_summed():
    parsed = parse_cv("""Sam Example
Experience
Analyst | January 2018 - January 2021
Consultant | January 2020 - January 2022
Education
Diploma
""")
    assert parsed.years_experience == 4.0


def test_present_uses_supplied_reference_date():
    parsed = parse_cv("""Alex Example
Experience
Reservation Agent | January 2020 - Present
""", reference_date=date(2025, 1, 1))
    assert parsed.years_experience == 5.0


def test_employment_dates_survive_pdf_column_heading_interleave():
    parsed = parse_cv("""Avery Example
Experience
Reservation Agent
Skills
Airline ticketing position | March 2016 - March 2020
Customer service role | April 2020 - April 2024
""", reference_date=date(2025, 1, 1))
    assert parsed.years_experience == 8.1


def test_split_date_endpoints_are_paired_and_overlaps_merged():
    parsed = parse_cv("""Jordan Example
Professional Experience
Travel Consultant
January 2015
December 2018
Reservation Officer
June 2018
June 2022
""", reference_date=date(2025, 1, 1))
    assert parsed.years_experience == 7.4


def test_arabic_month_and_now_variant_are_employment_dates():
    parsed = parse_cv("""ليلى مثال
الخبرة المهنية
موظفة حجوزات
مارس 2019
الآن
""", reference_date=date(2025, 3, 1))
    assert parsed.years_experience == 6.0


def test_non_employment_ranges_do_not_inflate_experience():
    parsed = parse_cv("""Taylor Example
Education
Bachelor degree | 2001 - 2005
Certifications
Professional course | 2010 - 2011
Projects
Migration project | 2012 - 2016
Personal details
Date of birth: 1980
""", reference_date=date(2025, 1, 1))
    assert parsed.years_experience is None


def test_current_job_does_not_pair_with_education_dates():
    parsed = parse_cv("""Morgan Example
Experience
Operations Manager | 2021 - Present
Education
University degree | 2003 - 2007
""", reference_date=date(2025, 1, 1))
    assert parsed.years_experience == 4.0


def test_single_adjacent_glyph_swap_in_long_arabic_skill_is_recovered():
    parsed = parse_cv("ليلى مثال\nالمهارات\nامايدوس\n" + "نص اصطناعي " * 30)
    assert "amadeus" in {skill.id for skill in parsed.skills}


def test_bilingual_docx_is_extracted_and_parsed_in_memory():
    document = Document()
    document.add_paragraph("Mona Example")
    document.add_paragraph("mona@example.test")
    document.add_paragraph("المهارات: A m a d e u s / N D C")
    document.add_paragraph("Education: Diploma")
    document.add_paragraph("Languages: Arabic, English")
    document.add_paragraph("Synthetic filler text. " * 15)
    buffer = BytesIO()
    document.save(buffer)
    extracted = extract_text(buffer.getvalue(), "synthetic.docx")
    parsed = parse_cv(extracted.text)
    assert not extracted.needs_manual_review
    assert {"amadeus", "ndc"} <= {skill.id for skill in parsed.skills}
    assert "NDC" in parsed.readable_text
    assert parsed.education == "diploma"


def test_empty_pdf_needs_manual_review():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)
    extracted = extract_text(buffer.getvalue(), "scan.pdf")
    assert extracted.pages == 1
    assert extracted.needs_manual_review


def test_size_limit_is_rejected():
    try:
        extract_text(b"x" * (5 * 1024 * 1024 + 1), "large.txt")
    except ValueError as exc:
        assert "5 MB" in str(exc)
    else:
        raise AssertionError("oversized CV was accepted")


def test_present_alias_is_matched_as_a_whole_word_only():
    parsed = parse_cv("""Jordan Example
Experience
Sales Agent | January 2018 - January 2022
Travel company representative at trade fairs
Skills
Communication & Presentations Skills
""", reference_date=date(2026, 9, 1))
    assert parsed.years_experience == 4.0


def test_arabic_professional_experience_heading_ends_contact_block():
    parsed = parse_cv("""سارة مثال
البيانات الشخصية
رقم الجوال: 0500000000
الملخص المهني
أخصائية خدمة عملاء
الخبرات المهنية
خدمة عملاء – شركة مثال (2024 - 2025)
• متابعة الحضور والانصراف.
""", reference_date=date(2026, 9, 1))
    assert parsed.years_experience == 1.0
    assert "experience" in parsed.sections


def test_open_range_rows_pair_with_their_own_end_despite_a_stray_year():
    parsed = parse_cv("""Robin Example
Experience
Travel Agent at Example Travel
From May 2008 to
May 2012
Award 2022
Senior Travel Consultant at Example Tours
From June 2019 to March
2022
Travel Consultant at Example Holidays
From January 2014 to
May 2016
""", reference_date=date(2026, 9, 1))
    # 4.0 + 2.75 + 2.33 years; the stray 2022 must not shift the pairing.
    assert parsed.years_experience == 9.1
