"""Integration coverage for the Phase 2 recruitment API (synthetic CVs only)."""
import base64
from itertools import count

from conftest import csrf
from app import db as appdb
from app.db import db


_seq = count(1)


def test_recruitment_schema_includes_semantic_phase_columns(client):
    with db() as conn:
        if appdb.USING_PG:
            tables = {
                row[0] for row in conn.execute(
                    "SELECT table_name FROM information_schema.tables WHERE table_schema=current_schema()"
                )
            }
            job_columns = {
                row[0] for row in conn.execute(
                    "SELECT column_name FROM information_schema.columns "
                    "WHERE table_schema=current_schema() AND table_name='job_openings'"
                )
            }
            applicant_columns = {
                row[0] for row in conn.execute(
                    "SELECT column_name FROM information_schema.columns "
                    "WHERE table_schema=current_schema() AND table_name='applicants'"
                )
            }
        else:
            tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            job_columns = {row[1] for row in conn.execute("PRAGMA table_info(job_openings)")}
            applicant_columns = {row[1] for row in conn.execute("PRAGMA table_info(applicants)")}
        assert {
            "job_openings", "applicants", "applicant_stage_history", "interview_feedback",
            "skills_taxonomy", "ats_semantic_runs", "ats_semantic_results",
        } <= tables
        assert {"rubric_version", "semantic_criteria"} <= job_columns
        assert {
            "cv_revision", "rules_score", "final_score", "semantic_status", "semantic_reason",
            "external_processing_consent_at", "consent_notice_version",
        } <= applicant_columns


def _cv(email: str, *, skills: str = "Amadeus, PNR, Excel, customer service", years: int = 4) -> str:
    # Long enough not to trigger the extractor's scanned/empty-document warning.
    text = f"""Jordan Synthetic
{email}
+966 55 000 1111
Professional Experience
Travel services specialist with {years} years experience handling reservations,
customer requests, airline ticket changes, reporting, and detailed client follow-up.
Education
Bachelor degree in a fictional discipline.
Languages
Arabic, English
Skills
{skills}
Additional profile
This entirely synthetic candidate maintained accurate records and collaborated
with a fictional operations team on service quality and process improvements.
"""
    return base64.b64encode(text.encode()).decode()


def _job_payload(title=None, department="Ops", required=None):
    return {
        "title": title or f"Synthetic ATS Job {next(_seq)}",
        "department": department,
        "location": "Riyadh",
        "employment_type": "full-time",
        "description": "Synthetic role used only by the automated test suite.",
        "min_years": 4,
        "min_education": "bachelor",
        "gates": {},
        "required_skills": required if required is not None else ["amadeus", "pnr"],
        "preferred_skills": ["excel", "customer_service"],
        "languages": ["Arabic", "English"],
        "weights": {
            "required_skills": 40,
            "experience": 20,
            "preferred_skills": 10,
            "education": 10,
            "languages": 10,
        },
        "semantic_criteria": [],
    }


def _create_job(client, **changes):
    payload = _job_payload()
    payload.update(changes)
    response = client.post("/api/recruitment/jobs", json=payload, headers=csrf(client))
    assert response.status_code == 201, response.text
    return response.json(), payload


def _upload(client, job_id, email=None, **changes):
    email = email or f"candidate{next(_seq)}@example.test"
    skills = changes.pop("skills", "Amadeus, PNR, Excel, customer service")
    years = changes.pop("years", 4)
    payload = {
        "source": "synthetic-test",
        "cv_attachment": _cv(email, skills=skills, years=years),
        "cv_name": "synthetic.txt",
        "cv_mime": "text/plain",
        "notice_attested": True,
        "notice_channel": "email",
    }
    payload.update(changes)
    return client.post(
        f"/api/recruitment/jobs/{job_id}/applicants", json=payload, headers=csrf(client)
    )


def test_role_matrix_and_department_scope(admin, new_employee):
    own_job, _ = _create_job(admin, department="Sales")
    other_job, _ = _create_job(admin, department="Finance")
    officer, _ = new_employee(role="hr_officer", department="Ops")
    head, _ = new_employee(role="dept_head", department="Sales")
    employee, _ = new_employee(role="employee", department="Sales")

    assert officer.get("/api/recruitment/jobs").status_code == 200
    assert officer.post(
        "/api/recruitment/jobs", json=_job_payload(department="Ops"), headers=csrf(officer)
    ).status_code == 201
    head_jobs = head.get("/api/recruitment/jobs")
    assert head_jobs.status_code == 200
    assert {job["department"] for job in head_jobs.json()} == {"Sales"}
    assert head.get(f"/api/recruitment/jobs/{own_job['id']}").status_code == 200
    assert head.get(f"/api/recruitment/jobs/{other_job['id']}").status_code == 403
    assert head.post(
        "/api/recruitment/jobs", json=_job_payload(), headers=csrf(head)
    ).status_code == 403
    assert employee.get("/api/recruitment/jobs").status_code == 403


def test_upload_score_round_trip_and_no_cv_leak(admin):
    job, _ = _create_job(admin)
    response = _upload(admin, job["id"])
    assert response.status_code == 201, response.text
    applicant = response.json()
    assert applicant["rules_score"] == applicant["score"] == applicant["final_score"]
    assert applicant["rules_score"] > 0
    assert applicant["semantic_status"] == "disabled"
    assert applicant["score_breakdown"]
    assert applicant["parsed"]["email"].endswith("@example.test")
    assert "cv_attachment" not in applicant and "cv_text" not in applicant

    listed = admin.get(f"/api/recruitment/jobs/{job['id']}/applicants").json()
    assert len(listed) == 1
    assert "cv_attachment" not in listed[0] and "cv_text" not in listed[0]
    assert "parsed" not in listed[0] and "corrections" not in listed[0]
    assert listed[0]["education"] == applicant["parsed"]["education"]
    assert listed[0]["years_experience"] == applicant["parsed"]["years_experience"]
    detail = admin.get(f"/api/recruitment/applicants/{applicant['id']}").json()
    assert "cv_attachment" not in detail and "cv_text" not in detail


def test_duplicate_email_is_scoped_to_job(admin):
    first_job, _ = _create_job(admin)
    second_job, _ = _create_job(admin)
    email = f"duplicate{next(_seq)}@example.test"
    assert _upload(admin, first_job["id"], email).status_code == 201
    duplicate = _upload(admin, first_job["id"], email)
    assert duplicate.status_code == 409
    assert duplicate.json()["detail"] == "duplicate_application"
    assert _upload(admin, second_job["id"], email).status_code == 201


def test_stage_transition_rules_and_history(admin):
    job, _ = _create_job(admin)
    applicant = _upload(admin, job["id"]).json()
    aid = applicant["id"]
    invalid = admin.post(
        f"/api/recruitment/applicants/{aid}/stage",
        json={"stage": "offer", "note": "skip"}, headers=csrf(admin),
    )
    assert invalid.status_code == 409
    moved = admin.post(
        f"/api/recruitment/applicants/{aid}/stage",
        json={"stage": "screened", "note": "reviewed"}, headers=csrf(admin),
    )
    assert moved.status_code == 200
    assert moved.json()["old_stage"] == "new" and moved.json()["new_stage"] == "screened"
    rejected = admin.post(
        f"/api/recruitment/applicants/{aid}/stage",
        json={"stage": "rejected", "note": "human decision"}, headers=csrf(admin),
    )
    assert rejected.status_code == 200
    assert admin.post(
        f"/api/recruitment/applicants/{aid}/stage",
        json={"stage": "screened"}, headers=csrf(admin),
    ).status_code == 200
    history = admin.get(f"/api/recruitment/applicants/{aid}").json()["history"]
    assert [(item["old_stage"], item["new_stage"]) for item in history] == [
        ("new", "screened"), ("screened", "rejected"), ("rejected", "screened")
    ]


def test_job_edit_and_correction_rescore(admin):
    job, payload = _create_job(admin, required_skills=["amadeus"])
    applicant = _upload(admin, job["id"], skills="Amadeus").json()
    original = applicant["rules_score"]

    payload["required_skills"] = ["sabre", "galileo"]
    edited = admin.put(
        f"/api/recruitment/jobs/{job['id']}", json=payload, headers=csrf(admin)
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["rescored_applicants"] == 1
    lower = admin.get(f"/api/recruitment/applicants/{applicant['id']}").json()["rules_score"]
    assert lower < original

    corrected = admin.patch(
        f"/api/recruitment/applicants/{applicant['id']}",
        json={"corrections": {"years_experience": 8, "education": "bachelor", "skills": ["sabre", "galileo"]}},
        headers=csrf(admin),
    )
    assert corrected.status_code == 200, corrected.text
    assert corrected.json()["rules_score"] > lower
    assert corrected.json()["corrections"]["skills"] == ["sabre", "galileo"]
    assert corrected.json()["parsed"]["skills"] != corrected.json()["corrections"]["skills"]


def test_cv_download_manager_only(admin, new_employee):
    job, _ = _create_job(admin)
    applicant = _upload(admin, job["id"]).json()
    downloaded = admin.get(f"/api/recruitment/applicants/{applicant['id']}/cv")
    assert downloaded.status_code == 200
    assert downloaded.json()["cv_name"] == "synthetic.txt"

    officer, _ = new_employee(role="hr_officer")
    assert officer.get(f"/api/recruitment/applicants/{applicant['id']}/cv").status_code == 403


def test_feedback_department_scope(admin, new_employee):
    job, _ = _create_job(admin, department="Operations")
    applicant = _upload(admin, job["id"]).json()
    own_head, _ = new_employee(role="dept_head", department="Operations")
    other_head, _ = new_employee(role="dept_head", department="Sales")
    response = own_head.post(
        f"/api/recruitment/applicants/{applicant['id']}/feedback",
        json={"result": "cleared", "ratings": [{"skill": "Service", "score_1_5": 4}]},
        headers=csrf(own_head),
    )
    assert response.status_code == 201
    assert response.json()["avg_rating"] == 4
    assert other_head.get(f"/api/recruitment/applicants/{applicant['id']}/feedback").status_code == 403


def test_taxonomy_seed_and_manager_edit(admin, new_employee):
    seeded = admin.get("/api/recruitment/taxonomy")
    assert seeded.status_code == 200
    assert any(item["id"] == "amadeus" for item in seeded.json())
    custom = {
        "id": f"synthetic_skill_{next(_seq)}",
        "canonical": "Synthetic Skill",
        "category": "test",
        "aliases": ["synthetic capability"],
        "is_active": True,
    }
    saved = admin.post("/api/recruitment/taxonomy", json=custom, headers=csrf(admin))
    assert saved.status_code == 200 and saved.json()["aliases"] == custom["aliases"]
    job, _ = _create_job(admin, required_skills=[custom["id"]], preferred_skills=[])
    matched = _upload(admin, job["id"], skills="synthetic capability").json()
    required = next(item for item in matched["score_breakdown"] if item["criterion"] == "required_skills")
    assert required["score_0_4"] == 4  # DB taxonomy edits affect future parsing/scoring.
    custom["canonical"] = "Updated Synthetic Skill"
    assert admin.post("/api/recruitment/taxonomy", json=custom, headers=csrf(admin)).json()["canonical"].startswith("Updated")
    officer, _ = new_employee(role="hr_officer")
    assert officer.get("/api/recruitment/taxonomy").status_code == 200
    assert officer.post("/api/recruitment/taxonomy", json=custom, headers=csrf(officer)).status_code == 403


def test_removed_alias_on_a_seeded_skill_stops_matching(admin):
    seeded = next(item for item in admin.get("/api/recruitment/taxonomy").json() if item["id"] == "bsp")
    job, _ = _create_job(admin, required_skills=["bsp"], preferred_skills=[])

    def required_score():
        applicant = _upload(admin, job["id"], skills="BSP reporting").json()
        return next(i for i in applicant["score_breakdown"] if i["criterion"] == "required_skills")["score_0_4"]

    assert required_score() == 4
    edited = {**seeded, "canonical": "Billing settlement", "aliases": ["billing and settlement plan"]}
    try:
        assert admin.post("/api/recruitment/taxonomy", json=edited, headers=csrf(admin)).status_code == 200
        assert required_score() == 0  # the seed's "bsp" alias must not come back
    finally:
        admin.post("/api/recruitment/taxonomy", json=seeded, headers=csrf(admin))

def test_parse_preview_does_not_persist_and_stats(admin):
    job, _ = _create_job(admin)
    before = admin.get("/api/recruitment/stats").json()["total_applicants"]
    email = f"preview{next(_seq)}@example.test"
    preview = admin.post(
        "/api/recruitment/parse-preview",
        json={"cv_attachment": _cv(email), "cv_name": "preview.txt", "cv_mime": "text/plain", "notice_attested": True},
        headers=csrf(admin),
    )
    assert preview.status_code == 200
    assert preview.json()["parsed"]["email"] == email
    assert admin.get("/api/recruitment/stats").json()["total_applicants"] == before

    applicant = _upload(admin, job["id"]).json()
    stats = admin.get("/api/recruitment/stats").json()
    assert stats["open_jobs"] >= 1
    assert stats["total_applicants"] == before + 1
    assert stats["applicants_per_stage"]["new"] >= 1
    job_stats = next(item for item in stats["jobs"] if item["job_id"] == job["id"])
    assert job_stats["applicant_count"] == 1
    assert job_stats["avg_score"] == applicant["score"]


def test_templates_list_for_hr_staff_only(admin, new_employee):
    response = admin.get("/api/recruitment/templates")
    assert response.status_code == 200
    templates = response.json()
    assert len(templates) == 6
    assert {"id", "title", "required_skills", "weights"} <= set(templates[0])
    head, _ = new_employee(role="dept_head", department="Sales")
    assert head.get("/api/recruitment/templates").status_code == 403


def test_export_ranks_applicants_and_neutralises_formulas(admin, new_employee):
    from io import BytesIO
    from openpyxl import load_workbook

    job, _ = _create_job(admin)
    assert _upload(admin, job["id"], full_name="=HYPERLINK(\"x\")").status_code == 201
    assert _upload(admin, job["id"], skills="Excel").status_code == 201
    response = admin.get(f"/api/recruitment/jobs/{job['id']}/export")
    assert response.status_code == 200
    rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
    assert rows[0][:3] == ("Rank", "ID", "Name")
    assert [row[0] for row in rows[1:]] == [1, 2]
    assert rows[1][9] >= rows[2][9]
    assert all(not str(row[2] or "").startswith("=") for row in rows[1:])
    head, _ = new_employee(role="dept_head", department="Ops")
    assert head.get(f"/api/recruitment/jobs/{job['id']}/export").status_code == 403
